import os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from PyPDF2 import PdfReader
import textract
import openpyxl
import docx

# Função pra extrair o conteudo
def extract_content(file_path):
  ext = file_path.split('.')[-1].lower()
  
  if ext == 'txt':
    with open(file_path, 'r') as file:
      content = file.read()
      
  elif ext == 'pdf':
    content = ''
    reader = PdfReader(file_path)
    for page in reader.pages:
      content += page.extract_text()
      
  elif ext == 'docx':
    doc = docx.Document(file_path)
    content = '\n'.join([para.text for para in doc.paragraphs])
    
  elif ext == 'xlsx':
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    content = ''
    for row in sheet.iter_rows(values_only=True):
      content += ' | '.join([str(cell) for cell in row]) + '\n'
  
  else:
    content = 'Arquivo não suportado para leitura.'
    
  return content

# Função pra gerar o relatorio em pdf
def generate_report(name, date, file_path, output_folder):
  if not name or not date:
    print('Nome ou Data não fornecidos!')
    return
  
  # extrair conteudo do arquivo selecionado
  file_content = extract_content(file_path)
  
  # caminho do pdf
  pdf_file = os.path.join(output_folder, 'relatorio.pdf')
  
  # cria o pdf
  c = canvas.Canvas(pdf_file, pagesize=letter)
  c.drawString(100, 750, f'Aqui está seu relatório, {name}')
  c.drawString(100, 730, f'Data: {date}')
  
  # adicionar mais conteudo ao pdf
  c.drawString(100, 710, 'Conteúdo do relatório aqui...')
  c.drawString(100, 690, file_content[:500]) # Limita a quantidade de texto mostrado
  
  # salva o pdf
  c.save()
  print(f'Relatório gerado com sucesso: {pdf_file}')